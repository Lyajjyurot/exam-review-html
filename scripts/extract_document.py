#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
extract_document.py —— 把 PDF / Word / 纯文本提纲抽取为结构化 Markdown，供后续批注与出题使用。

用法:
    python extract_document.py <文件或目录> [...] -o <输出目录> [--json] [--max-pages N]

输出:
    <输出目录>/<原文件名>.md      带层级标题与分页标记的文本
    <输出目录>/_manifest.json     (加 --json 时) 每个文件的页数/字数/标题数统计

依赖策略:
    .docx  -> 优先 python-docx；缺失时用标准库 zipfile 解析（保留标题层级/表格/加粗）
    .pdf   -> 依次尝试 pypdf / pdfplumber / PyMuPDF；全部缺失时提示安装命令
    .txt/.md -> 原样读取，仅做换行归一化
    .doc   -> 老二进制格式，脚本无法解析，提示转换为 .docx 后重试

Windows 下请确保 stdout 编码为 UTF-8（脚本内部已强制重配置）。
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:  # pragma: no cover
    pass

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".markdown", ".doc"}

# --------------------------------------------------------------------------- #
# Word (.docx)
# --------------------------------------------------------------------------- #


def _docx_with_library(path: str) -> str:
    import docx  # type: ignore

    doc = docx.Document(path)
    out: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading") or re.match(r"^标题\s*\d", style) or style.startswith("title"):
            m = re.search(r"(\d+)", style)
            level = int(m.group(1)) if m else 2
            out.append("#" * min(max(level, 1), 6) + " " + text)
        elif re.match(r"^\d+(\.\d+)*[、.\s]", text) and len(text) < 60:
            out.append("### " + text)
        else:
            out.append(text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
        if rows:
            out.append("")
            out.append(rows[0])
            out.append("| " + " | ".join("---" for _ in rows[0].split("|")[1:-1]) + " |")
            out.extend(rows[1:])
            out.append("")
    return "\n\n".join(out)


def _para_text(node: ET.Element) -> tuple[str, bool]:
    """返回 (文本, 是否整段加粗)。"""
    buf: list[str] = []
    bold = True
    has_run = False
    for t in node.iter():
        if t.tag == W_NS + "t":
            buf.append(t.text or "")
        elif t.tag == W_NS + "br":
            buf.append("\n")
    for r in node.findall(W_NS + "r"):
        has_run = True
        rpr = r.find(W_NS + "rPr")
        if rpr is None or rpr.find(W_NS + "b") is None:
            bold = False
    return "".join(buf).strip(), bool(bold and has_run and buf)


def _docx_with_stdlib(path: str) -> str:
    out: list[str] = []
    with zipfile.ZipFile(path) as z:
        xml_bytes = z.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    body = root.find(W_NS + "body")
    if body is None:
        return ""

    for child in body:
        if child.tag == W_NS + "p":
            text, bold = _para_text(child)
            if not text:
                continue
            ppr = child.find(W_NS + "pPr")
            style_id = ""
            outline = ""
            if ppr is not None:
                ps = ppr.find(W_NS + "pStyle")
                if ps is not None:
                    style_id = (ps.get(W_NS + "val") or "").lower()
                ol = ppr.find(W_NS + "outlineLvl")
                if ol is not None:
                    outline = ol.get(W_NS + "val") or ""
            level = None
            m = re.search(r"heading\s*(\d)", style_id) or re.search(r"(\d)", style_id)
            if "heading" in style_id or style_id.startswith("1") or "标题" in style_id:
                level = int(m.group(1)) if m else 2
            if level is None and outline:
                level = int(outline) + 1
            if level is None and re.match(r"^第?[一二三四五六七八九十\d]+[章节篇部分]", text) and len(text) < 40:
                level = 2
            if level is None and re.match(r"^\d+(\.\d+)*[、.\s]", text) and len(text) < 60:
                level = 4
            if level:
                out.append("#" * min(max(level, 1), 6) + " " + text)
            elif bold and len(text) < 60:
                out.append("**" + text + "**")
            else:
                out.append(text)

        elif child.tag == W_NS + "tbl":
            rows = []
            for tr in child.findall(W_NS + "tr"):
                cells = []
                for tc in tr.findall(W_NS + "tc"):
                    txt = " ".join(
                        (t.text or "") for p in tc.iter(W_NS + "p") for t in p.iter(W_NS + "t")
                    ).strip()
                    cells.append(txt.replace("\n", " "))
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                out.append("")
                out.append(rows[0])
                out.append("| " + " | ".join("---" for _ in rows[0].split("|")[1:-1]) + " |")
                out.extend(rows[1:])
                out.append("")
    return "\n\n".join(out)


def extract_docx(path: str) -> str:
    try:
        text = _docx_with_library(path)
        if text.strip():
            return text
    except ImportError:
        pass
    except Exception as e:
        print(f"[warn] python-docx 解析失败({e})，回退到标准库解析", file=sys.stderr)
    return _docx_with_stdlib(path)


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def _pdf_outline(reader) -> list[str]:
    lines: list[str] = []
    try:
        outlines = reader.outline or []

        def walk(items, depth=0):
            for it in items:
                if isinstance(it, list):
                    walk(it, depth + 1)
                else:
                    lines.append("  " * depth + "- " + str(getattr(it, "title", it)))

        walk(outlines)
    except Exception:
        pass
    return lines


def extract_pdf(path: str, max_pages: int | None = None) -> str:
    # 1) pypdf
    try:
        import pypdf  # type: ignore

        reader = pypdf.PdfReader(path)
        total = len(reader.pages)
        limit = min(total, max_pages) if max_pages else total
        chunks = []
        outline = _pdf_outline(reader)
        if outline:
            chunks.append("<!-- PDF 书签目录 -->\n" + "\n".join(outline) + "\n")
        for i in range(limit):
            txt = reader.pages[i].extract_text() or ""
            chunks.append(f"\n<!-- page {i + 1} -->\n" + txt.strip())
        if max_pages and total > max_pages:
            chunks.append(f"\n<!-- 已截断：原文件 {total} 页，仅抽取前 {max_pages} 页 -->")
        return "".join(chunks)
    except ImportError:
        pass

    # 2) pdfplumber
    try:
        import pdfplumber  # type: ignore

        chunks = []
        with pdfplumber.open(path) as pdf:
            total = len(pdf.pages)
            limit = min(total, max_pages) if max_pages else total
            for i, page in enumerate(pdf.pages[:limit]):
                chunks.append(f"\n<!-- page {i + 1} -->\n" + (page.extract_text() or "").strip())
            if max_pages and total > max_pages:
                chunks.append(f"\n<!-- 已截断：原文件 {total} 页，仅抽取前 {max_pages} 页 -->")
        return "".join(chunks)
    except ImportError:
        pass

    # 3) PyMuPDF
    try:
        import fitz  # type: ignore

        doc = fitz.open(path)
        total = doc.page_count
        limit = min(total, max_pages) if max_pages else total
        chunks = []
        for i in range(limit):
            chunks.append(f"\n<!-- page {i + 1} -->\n" + doc[i].get_text("text").strip())
        if max_pages and total > max_pages:
            chunks.append(f"\n<!-- 已截断：原文件 {total} 页，仅抽取前 {max_pages} 页 -->")
        return "".join(chunks)
    except ImportError:
        pass

    raise RuntimeError(
        "未找到 PDF 解析库。请在隔离虚拟环境中安装后重试：\n"
        "  <managed-python> -m pip install pypdf\n"
        "若网络受限，可改用国内镜像： -i https://pypi.tuna.tsinghua.edu.cn/simple"
    )


# --------------------------------------------------------------------------- #
# 通用
# --------------------------------------------------------------------------- #


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    # 去掉 PDF 常见的页眉页脚孤立页码行
    text = re.sub(r"\n\s*\d{1,3}\s*(?=\n)", "\n", text)
    return text.strip()


def extract(path: str, max_pages: int | None = None) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return normalize(extract_docx(path))
    if ext == ".pdf":
        return normalize(extract_pdf(path, max_pages))
    if ext in {".txt", ".md", ".markdown"}:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return normalize(f.read())
    if ext == ".doc":
        raise RuntimeError(f"{os.path.basename(path)} 是旧版 .doc 二进制格式，请先另存为 .docx 再重试。")
    raise RuntimeError(f"暂不支持的文件类型: {ext}")


def stats(text: str) -> dict:
    headings = re.findall(r"^#{1,6}\s+.+$", text, flags=re.M)
    return {
        "chars": len(text),
        "headings": len(headings),
        "pages": len(re.findall(r"<!-- page \d+ -->", text)),
        "outline": [h.lstrip("# ").strip() for h in headings[:40]],
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="抽取 PDF/Word 提纲为 Markdown 文本")
    ap.add_argument("inputs", nargs="+", help="文件或目录")
    ap.add_argument("-o", "--out", default="./extracted", help="输出目录，默认 ./extracted")
    ap.add_argument("--json", action="store_true", help="额外输出 _manifest.json")
    ap.add_argument("--max-pages", type=int, default=None, help="PDF 最多抽取页数")
    args = ap.parse_args()

    os.makedirs(args.out, exist_ok=True)

    files: list[str] = []
    for p in args.inputs:
        if os.path.isdir(p):
            for root, _, names in os.walk(p):
                for n in names:
                    if os.path.splitext(n)[1].lower() in SUPPORTED and not n.startswith("~$"):
                        files.append(os.path.join(root, n))
        else:
            files.append(p)

    if not files:
        print("未找到可解析的文件（支持 .pdf / .docx / .txt / .md）")
        return 1

    manifest = []
    for f in files:
        name = os.path.splitext(os.path.basename(f))[0]
        try:
            text = extract(f, args.max_pages)
        except Exception as e:
            print(f"[error] {name}: {e}")
            manifest.append({"file": f, "error": str(e)})
            continue
        out_path = os.path.join(args.out, name + ".md")
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(text)
        s = stats(text)
        s["file"] = f
        s["output"] = out_path
        manifest.append(s)
        print(f"[ok] {name:<28} 字数 {s['chars']:>6}  标题 {s['headings']:>3}  页 {s['pages']:>3}  -> {out_path}")

    if args.json:
        import json

        with open(os.path.join(args.out, "_manifest.json"), "w", encoding="utf-8") as fh:
            json.dump(manifest, fh, ensure_ascii=False, indent=2)

    bad = [m for m in manifest if "error" in m]
    if bad:
        print(f"\n{len(bad)} 个文件解析失败，请查看上方 [error] 提示。")
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
